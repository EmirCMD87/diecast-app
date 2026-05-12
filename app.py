from flask import Flask, render_template, request, redirect, url_for, flash, send_from_directory, send_file, session
from flask_sqlalchemy import SQLAlchemy
from flask_login import LoginManager, UserMixin, login_user, login_required, logout_user, current_user
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from datetime import datetime, timedelta
from functools import wraps
import os
import tempfile
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

app = Flask(__name__)
app.config['UPLOAD_FOLDER'] = 'uploads'
app.config['SEND_FILE_MAX_AGE_DEFAULT'] = 0
app.secret_key = "diecast_gizli_anahtar_123"
ADMIN_PASSWORD = "diecast_admin_2025"

# Oturum ayarları
app.config['REMEMBER_COOKIE_DURATION'] = timedelta(days=365)
app.config['REMEMBER_COOKIE_HTTPONLY'] = True
app.config['REMEMBER_COOKIE_SECURE'] = True
app.config['SESSION_PERMANENT'] = True
app.config['PERMANENT_SESSION_LIFETIME'] = timedelta(days=365)

# Veritabanı
app.config["SQLALCHEMY_DATABASE_URI"] = "sqlite:///diecast.db"
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False

# Dosya yükleme
UPLOAD_FOLDER = "uploads"
ALLOWED_EXTENSIONS = {"png", "jpg", "jpeg", "gif", "webp"}
app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

db = SQLAlchemy(app)
login_manager = LoginManager(app)
login_manager.login_view = "login"

# ------------------- VERİTABANI MODELLERİ -------------------
class User(UserMixin, db.Model):
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(80), unique=True, nullable=False)
    password = db.Column(db.String(200), nullable=False)
    arabalar = db.relationship("Araba", backref="sahip", lazy=True)

class Araba(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    isim = db.Column(db.String(100), nullable=False)
    marka = db.Column(db.String(50), nullable=False)
    renk = db.Column(db.String(50), nullable=False)
    olcek = db.Column(db.String(20), nullable=False, default="1:64")
    resim_yolu = db.Column(db.String(200), nullable=False)
    tarih = db.Column(db.DateTime, default=datetime.utcnow)
    user_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False)

# ------------------- YARDIMCI -------------------
def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS

@login_manager.user_loader
def load_user(user_id):
    return User.query.get(int(user_id))

def admin_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if current_user.username != "EmirCMD87":
            flash("Bu sayfaya erişim yetkin yok", "danger")
            return redirect(url_for("dashboard"))
        if not session.get('admin_verified'):
            return redirect(url_for('admin_login'))
        return f(*args, **kwargs)
    return decorated_function

@app.route('/uploads/<filename>')
def uploaded_file(filename):
    return send_from_directory('uploads', filename)

# ------------------- ROTALAR -------------------
@app.route("/")
def index():
    return redirect(url_for("login"))

@app.route("/register", methods=["GET", "POST"])
def register():
    if request.method == "POST":
        username = request.form["username"].strip()
        password = request.form["password"]
        
        if not username or not password:
            flash("Kullanıcı adı ve şifre boş olamaz!", "danger")
            return redirect(url_for("register"))
        
        existing_user = User.query.filter_by(username=username).first()
        if existing_user:
            flash("❌ Bu kullanıcı adı zaten alınmış!", "danger")
            return redirect(url_for("register"))
        
        hashed_password = generate_password_hash(password, method="pbkdf2:sha256")
        new_user = User(username=username, password=hashed_password)
        db.session.add(new_user)
        db.session.commit()
        flash("✅ Kayıt başarılı! Şimdi giriş yapabilirsin.", "success")
        return redirect(url_for("login"))
    
    return render_template("register.html")

@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username = request.form["username"].strip()
        password = request.form["password"]
        user = User.query.filter_by(username=username).first()
        if user and check_password_hash(user.password, password):
            remember = True if request.form.get('remember') else False
            login_user(user, remember=remember)
            flash(f"Hoş geldin, {username}! ✅", "success")
            return redirect(url_for("dashboard"))
        else:
            flash("❌ Kullanıcı adı veya şifre hatalı!", "danger")
    return render_template("login.html")

@app.route("/logout")
@login_required
def logout():
    logout_user()
    session.pop('admin_verified', None)
    flash("Çıkış yapıldı. 👋", "info")
    return redirect(url_for("login"))

@app.route("/dashboard")
@login_required
def dashboard():
    marka_filtre = request.args.get('marka', '')
    renk_filtre = request.args.get('renk', '')
    olcek_filtre = request.args.get('olcek', '')
    arama = request.args.get('arama', '')
    
    query = Araba.query.filter_by(user_id=current_user.id)
    if marka_filtre:
        query = query.filter_by(marka=marka_filtre)
    if renk_filtre:
        query = query.filter_by(renk=renk_filtre)
    if olcek_filtre:
        query = query.filter_by(olcek=olcek_filtre)
    if arama:
        query = query.filter(Araba.isim.contains(arama))
    arabalar = query.order_by(Araba.tarih.desc()).all()
    
    tum_arabalar = Araba.query.filter_by(user_id=current_user.id).all()
    toplam_araba = len(tum_arabalar)
    renk_sayilari = {}
    marka_sayilari = {}
    olcek_sayilari = {}
    for a in tum_arabalar:
        renk_sayilari[a.renk] = renk_sayilari.get(a.renk, 0) + 1
        marka_sayilari[a.marka] = marka_sayilari.get(a.marka, 0) + 1
        olcek_sayilari[a.olcek] = olcek_sayilari.get(a.olcek, 0) + 1
    
    return render_template("dashboard.html",
                         arabalar=arabalar,
                         toplam_araba=toplam_araba,
                         renk_sayilari=renk_sayilari,
                         marka_sayilari=marka_sayilari,
                         olcek_sayilari=olcek_sayilari,
                         secili_marka=marka_filtre,
                         secili_renk=renk_filtre,
                         secili_olcek=olcek_filtre,
                         arama_kelimesi=arama)

@app.route("/araba_ekle", methods=["GET", "POST"])
@login_required
def araba_ekle():
    if request.method == "POST":
        isim = request.form["isim"]
        marka = request.form["marka"]
        renk = request.form["renk"]
        olcek = request.form["olcek"]
        dosya = request.files["resim"]
        
        if not dosya or dosya.filename == "":
            flash("Resim seçmediniz!", "danger")
            return redirect(request.url)
        if not allowed_file(dosya.filename):
            flash("Sadece resim dosyaları yüklenebilir!", "danger")
            return redirect(request.url)
        
        filename = secure_filename(f"{current_user.id}_{datetime.now().timestamp()}_{dosya.filename}")
        dosya_yolu = os.path.join(app.config["UPLOAD_FOLDER"], filename)
        dosya.save(dosya_yolu)
        
        yeni_araba = Araba(
            isim=isim, 
            marka=marka, 
            renk=renk, 
            olcek=olcek,
            resim_yolu=dosya_yolu, 
            user_id=current_user.id
        )
        db.session.add(yeni_araba)
        db.session.commit()
        flash(f"{isim} başarıyla eklendi!", "success")
        return redirect(url_for("dashboard"))
    
    return render_template("araba_ekle.html")

@app.route("/araba_sil/<int:araba_id>")
@login_required
def araba_sil(araba_id):
    araba = Araba.query.get_or_404(araba_id)
    if araba.user_id != current_user.id:
        flash("Bu arabayı silme yetkiniz yok!", "danger")
        return redirect(url_for("dashboard"))
    if os.path.exists(araba.resim_yolu):
        os.remove(araba.resim_yolu)
    db.session.delete(araba)
    db.session.commit()
    flash(f"{araba.isim} silindi.", "info")
    return redirect(url_for("dashboard"))

@app.route("/export_excel")
@login_required
def export_excel():
    arabalar = Araba.query.filter_by(user_id=current_user.id).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Koleksiyonum"
    basliklar = ["ID", "Araba Adı", "Marka", "Renk", "Ölçek", "Eklenme Tarihi"]
    for col, baslik in enumerate(basliklar, 1):
        cell = ws.cell(row=1, column=col, value=baslik)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1E3C72", end_color="1E3C72", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")
    for row, a in enumerate(arabalar, 2):
        ws.cell(row=row, column=1, value=a.id)
        ws.cell(row=row, column=2, value=a.isim)
        ws.cell(row=row, column=3, value=a.marka)
        ws.cell(row=row, column=4, value=a.renk)
        ws.cell(row=row, column=5, value=a.olcek)
        ws.cell(row=row, column=6, value=a.tarih.strftime('%Y-%m-%d %H:%M'))
    for col in range(1, 7):
        ws.column_dimensions[get_column_letter(col)].width = 20
    temp = tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx')
    wb.save(temp.name)
    temp.close()
    return send_file(temp.name, as_attachment=True, download_name=f"koleksiyonum_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx", mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.route('/@<username>')
def public_profile(username):
    user = User.query.filter_by(username=username).first_or_404()
    arabalar = Araba.query.filter_by(user_id=user.id).order_by(Araba.tarih.desc()).all()
    return render_template('public_profile.html', user=user, arabalar=arabalar)

# ------------------- ADMIN PANELİ -------------------
@app.route("/admin/login", methods=["GET", "POST"])
def admin_login():
    if request.method == "POST":
        pwd = request.form.get("password")
        if pwd == ADMIN_PASSWORD:
            session['admin_verified'] = True
            flash("Admin girişi başarılı!", "success")
            return redirect(url_for("admin_panel"))
        else:
            flash("Şifre yanlış!", "danger")
    return render_template("admin_login.html")

@app.route("/admin")
@login_required
@admin_required
def admin_panel():
    kullanicilar = User.query.all()
    return render_template("admin.html", kullanicilar=kullanicilar)

# ------------------- VERİTABANI OLUŞTUR -------------------
with app.app_context():
    db.create_all()
    print("✅ Veritabanı hazır.")

if __name__ == "__main__":
    app.run(debug=True, host="0.0.0.0", port=5000)
